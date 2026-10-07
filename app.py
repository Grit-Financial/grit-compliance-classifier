from io import BytesIO
from datetime import datetime
import re
import uuid

import httpx
import pandas as pd
from fastapi import FastAPI, HTTPException
from fastapi.responses import HTMLResponse, Response
from pydantic import BaseModel
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app = FastAPI(title="Grit Compliance Complaint Classifier")
RUNS = {}

SHEET_RE = re.compile(r"/spreadsheets/d/([a-zA-Z0-9-_]+)")

REGS = {
    "UDAAP": {
        "strong": ["mislead","not disclosed","undisclosed","decept","unable to access","cannot access","unexpected fee","promised"],
        "medium": ["confus","fee","locked","frozen","declin","limit","delay","missing funds","charged"]
    },
    "Reg E": {
        "strong": ["unauthorized","did not authorize","dispute","missing funds","atm","debit card","ach","electronic transfer"],
        "medium": ["transfer","withdraw","card","reversal","transaction","deposit"]
    },
    "Reg DD": {
        "strong": ["apy","annual percentage yield","deposit disclosure","interest rate"],
        "medium": ["deposit account","balance requirement","account fee"]
    },
    "Reg P": {
        "strong": ["privacy","personal information","nonpublic personal","shared my data","data shared"],
        "medium": ["third party","information disclosed","data"]
    },
    "Reg GG": {
        "strong": ["gambling","wager","betting","casino"],
        "medium": ["restricted transaction"]
    }
}

CATEGORIES = {
    "Account Opening Issue":["open account","onboarding","identity verification","kyc","registration"],
    "Account Restriction":["locked","frozen","restricted","suspended","disabled","blocked"],
    "Customer Service":["customer service","support","no response","callback","follow up"],
    "Deposit/Withdrawal Issue":["deposit","withdraw","withdrawal","missing funds","bank transfer"],
    "Fees":["fee","charge","charged"],
    "Unresolved Fraud Non ID Theft":["unauthorized","fraud","not mine","did not authorize","dispute"],
    "Unresolved ID Theft":["identity theft","stolen identity"],
    "Interest":["apy","interest","yield"],
    "Limits":["limit","maximum","cap","exceeded"],
    "Marketing/Advertising":["advertis","promotion","promised","marketing","misleading"],
    "Return of Funds":["refund","reversal","reversed","credited back"],
    "Statement Issue":["statement","transaction history"],
    "Technical":["bug","error","app","system","crash","technical","failed"]
}

ROOTS = {
    "EWA Availability / Limit Issue":["ewa","earned wage","available wage","accrued wage","repayment","limit"],
    "Identity / KYC Verification Issue":["kyc","identity","verify","verification","phone number","address verification"],
    "Onboarding / Account Setup Issue":["onboard","registration","sign up","account opening"],
    "Deposit / Withdrawal / Return of Funds Issue":["deposit","withdraw","refund","reversal","bank transfer","funds"],
    "System / Technical Issue":["technical","system","bug","app","failed","error"],
    "Card / ATM / Transaction Decline Issue":["card","atm","declined","pin","merchant","fraud detection"],
    "Customer Education / Disclosure Clarity Issue":["explained","informed","advised","fee","policy","confus"],
    "Unauthorized Transaction / Dispute Issue":["unauthorized","dispute","fraud","not mine"]
}

class AnalyzeRequest(BaseModel):
    sheet_url: str

def contains(text, term):
    return term.lower() in text.lower()

def pick_label(text, mapping, fallback):
    scored = [(sum(1 for t in terms if contains(text,t)), label) for label,terms in mapping.items()]
    count,label = max(scored)
    return label if count else fallback

def score_reg(text, cfg):
    strong = [x for x in cfg["strong"] if contains(text,x)]
    medium = [x for x in cfg["medium"] if contains(text,x)]
    if strong:
        score = min(98, 76 + 9*len(strong) + min(6,3*len(medium)))
    elif medium:
        score = min(86, 55 + 8*len(medium))
    else:
        score = 0
    return score,strong+medium

def classify(rec):
    complaint = str(rec.get("complaint","") or "")
    resolution = str(rec.get("resolution","") or "")
    text = (complaint+" "+resolution).lower()
    scores, hits = {}, {}
    for reg,cfg in REGS.items():
        scores[reg],hits[reg] = score_reg(text,cfg)
    potential = [r for r,s in scores.items() if s >= 65]
    top = max(scores.values()) if scores else 0

    if potential:
        confidence = max(top, min(96,70+4*len(potential)))
        proposed = " + ".join(potential)
    else:
        wc = len(text.split())
        confidence = 82 if wc >= 18 else 58 if wc >= 8 else 35
        proposed = "Operational / No Clear Regulatory Indicator"

    if confidence < 50:
        primary = "Unknown / Insufficient Evidence"
        review = "Yes"
    elif confidence < 80:
        primary = "Needs Human Review"
        review = "Yes"
    else:
        primary = proposed
        review = "No"

    severity = 5 if any(x in text for x in ["unauthorized","fraud","identity theft","missing funds"]) else 4 if any(x in text for x in ["locked","frozen","cannot access","unable to access","refund","reversal"]) else 3 if any(x in text for x in ["fee","declined","limit","delay"]) else 2 if len(text.split())>10 else 1
    likelihood = 3 if top>=90 else 2 if top>=75 else 1 if top>=60 else 0
    priority = "High" if severity>=4 and likelihood>=2 else "Medium" if severity>=3 or review=="Yes" else "Standard"

    trigger_text = []
    for reg in potential:
        if hits[reg]:
            trigger_text.append(reg+": "+", ".join(sorted(set(hits[reg]))[:4]))
    if not trigger_text:
        trigger_text = ["No clear regulatory trigger identified from the available narrative"]

    missing = []
    if "Reg E" in potential:
        if "$" not in text: missing.append("transaction amount")
        if not any(x in text for x in ["unauthorized","authorize","fraud","dispute"]): missing.append("authorization status")
    if "UDAAP" in potential and not any(x in text for x in ["disclos","explain","inform","promis","mislead"]):
        missing.append("relevant disclosure or customer communication")

    rationale = []
    for reg in potential:
        rationale.append(f"{reg} indicator supported by complaint facts and trigger terms; confirm applicability against the complete record.")
    if not rationale:
        rationale.append("No clear UDAAP, Regulation E, Regulation DD, Regulation P, or Regulation GG trigger was identified from the available complaint narrative.")
    if review == "Yes":
        rationale.append("Human review is required under the confidence policy.")

    return {
        **rec,
        "issue_category": pick_label(text,CATEGORIES,"Other"),
        "root_cause_category": pick_label(text,ROOTS,"Unable to Determine from Complaint Log"),
        "udaap":"Yes" if "UDAAP" in potential else "No",
        "reg_e":"Yes" if "Reg E" in potential else "No",
        "reg_dd":"Yes" if "Reg DD" in potential else "No",
        "reg_p":"Yes" if "Reg P" in potential else "No",
        "reg_gg":"Yes" if "Reg GG" in potential else "No",
        "potential_regulations":", ".join(potential) if potential else "None identified",
        "primary_classification":primary,
        "proposed_classification":proposed,
        "confidence":round(float(confidence),1),
        "needs_human_review":review,
        "triggering_facts":"; ".join(trigger_text),
        "missing_facts":", ".join(missing) if missing else "None identified from current rule set",
        "regulatory_likelihood":likelihood,
        "consumer_impact_severity":severity,
        "priority":priority,
        "recommended_action":"Compliance review required" if review=="Yes" else ("Compliance review recommended" if priority=="High" else "Standard compliance triage"),
        "regulatory_rationale":" ".join(rationale),
        "classification_basis":"Deterministic regulatory rules + contextual heuristics",
        "rule_version":"2026.10-mvp.1"
    }

def get_sheet(url):
    m = SHEET_RE.search(url)
    if not m: raise ValueError("Please provide a valid Google Sheets URL.")
    sid = m.group(1)
    gm = re.search(r"(?:gid=|#gid=)(\d+)",url)
    gid = gm.group(1) if gm else "0"
    export = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=csv&gid={gid}"
    r = httpx.get(export,follow_redirects=True,timeout=30)
    if r.status_code != 200 or "text/html" in r.headers.get("content-type",""):
        raise PermissionError("The sheet could not be read. For this MVP, the sheet must be readable by the application. Private Workspace access will be added before production use.")
    return pd.read_csv(BytesIO(r.content))

def infer(df):
    low = {str(c).strip().lower():c for c in df.columns}
    def pick(words,required=False):
        for w in words:
            for k,v in low.items():
                if w == k or w in k: return v
        if required: raise ValueError("Could not find a complaint/reason/description/narrative column.")
        return ""
    return {
        "complaint":pick(["complaint","reason","description","narrative","issue"],True),
        "resolution":pick(["resolution","response","outcome","action taken"]),
        "case_id":pick(["case id","ticket","complaint id","id"]),
        "date":pick(["complaint date","date","created"]),
        "customer":pick(["customer","name","employee"]),
        "program":pick(["program","product","client"]),
        "source":pick(["source","channel","method"])
    }

def safe(v):
    return "" if pd.isna(v) else str(v)

def summarize(rows):
    n=len(rows)
    return {
        "total_cases":n,
        "needs_human_review":sum(r["needs_human_review"]=="Yes" for r in rows),
        "average_confidence":round(sum(r["confidence"] for r in rows)/n,1) if n else 0,
        "udaap_flags":sum(r["udaap"]=="Yes" for r in rows),
        "reg_e_flags":sum(r["reg_e"]=="Yes" for r in rows),
        "reg_dd_flags":sum(r["reg_dd"]=="Yes" for r in rows),
        "reg_p_flags":sum(r["reg_p"]=="Yes" for r in rows),
        "reg_gg_flags":sum(r["reg_gg"]=="Yes" for r in rows),
        "unknown":sum(r["primary_classification"]=="Unknown / Insufficient Evidence" for r in rows),
        "high_priority":sum(r["priority"]=="High" for r in rows)
    }

def workbook(rows,summary):
    out=BytesIO()
    df=pd.DataFrame(rows)
    review=df[df["needs_human_review"]=="Yes"].copy() if not df.empty else df.copy()
    summ=pd.DataFrame([{"Metric":k.replace("_"," ").title(),"Value":v} for k,v in summary.items()])
    with pd.ExcelWriter(out,engine="openpyxl") as writer:
        df.to_excel(writer,index=False,sheet_name="Classification Results")
        review.to_excel(writer,index=False,sheet_name="Human Review")
        summ.to_excel(writer,index=False,sheet_name="Run Summary")
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:
                c.fill=PatternFill("solid",fgColor="35123F"); c.font=Font(color="FFFFFF",bold=True); c.alignment=Alignment(vertical="center")
            for i,col in enumerate(ws.columns,1):
                width=min(60,max((len(str(c.value)) if c.value is not None else 0 for c in col),default=8)+2)
                ws.column_dimensions[get_column_letter(i)].width=max(10,width)
    return out.getvalue()

@app.get("/",response_class=HTMLResponse)
def home():
    return HTML

@app.get("/api/health")
def health():
    return {"status":"ok"}

@app.post("/api/analyze")
def analyze(req:AnalyzeRequest):
    try:
        df=get_sheet(req.sheet_url)
        cols=infer(df)
        rows=[]
        for idx,row in df.iterrows():
            rec={
                "row_number":int(idx)+2,
                "case_id":safe(row[cols["case_id"]]) if cols["case_id"] else f"ROW-{idx+2}",
                "customer":safe(row[cols["customer"]]) if cols["customer"] else "",
                "date":safe(row[cols["date"]]) if cols["date"] else "",
                "program_product":safe(row[cols["program"]]) if cols["program"] else "",
                "source_channel":safe(row[cols["source"]]) if cols["source"] else "",
                "complaint":safe(row[cols["complaint"]]),
                "resolution":safe(row[cols["resolution"]]) if cols["resolution"] else ""
            }
            rows.append(classify(rec))
        run_id=str(uuid.uuid4())
        summary=summarize(rows)
        RUNS[run_id]={"rows":rows,"summary":summary}
        return {"run_id":run_id,"summary":summary,"results":rows}
    except (ValueError,PermissionError) as e:
        raise HTTPException(status_code=400,detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500,detail=f"Analysis failed: {e}")

@app.get("/api/runs/{run_id}/download.xlsx")
def download(run_id:str):
    run=RUNS.get(run_id)
    if not run: raise HTTPException(status_code=404,detail="Run not found or expired.")
    content=workbook(run["rows"],run["summary"])
    name=f"Grit_Complaint_Classification_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(content=content,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{name}"'})

HTML = r'''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Grit Compliance Analytics</title><style>
:root{--pink:#ee3296;--fuchsia:#c814c7;--purple:#8500e6;--plum:#35123f;--ink:#171218;--muted:#6d6e71;--line:#eadfea;--red:#b42318;--green:#166534;--amber:#a15c00;--grad:linear-gradient(90deg,var(--pink),var(--fuchsia) 52%,var(--purple));--shadow:0 14px 38px rgba(53,18,63,.09)}*{box-sizing:border-box}body{margin:0;font-family:Inter,system-ui,sans-serif;color:var(--ink);background:linear-gradient(180deg,#fff,#fbf9fc 30rem,#f8f5f9)}.rule{height:6px;background:var(--grad)}header{background:#fff;border-bottom:1px solid var(--line);padding:18px 30px 28px}.top{max-width:1600px;margin:auto;display:flex;justify-content:space-between;align-items:center}.brand{font-size:24px;font-weight:900}.brand span{color:var(--purple)}.status{font-size:12px;color:var(--muted)}.hero{max-width:1600px;margin:18px auto 0;background:var(--grad);color:#fff;border-radius:22px;padding:28px 30px}.hero h1{font-size:clamp(28px,3.2vw,46px);margin:5px 0 10px;line-height:1.05}.hero p{max-width:1100px;margin:0;opacity:.92}.eyebrow{text-transform:uppercase;letter-spacing:.12em;font-size:11px;font-weight:850}main{max-width:1600px;margin:auto;padding:22px 28px 55px}.notice,.panel{background:#fff;border:1px solid var(--line);border-radius:15px;box-shadow:var(--shadow)}.notice{border-left:5px solid var(--fuchsia);padding:15px 17px;font-size:13px;color:#514957}.panel{margin:18px 0;overflow:hidden}.head{padding:15px 17px;border-bottom:1px solid var(--line);display:flex;justify-content:space-between;align-items:center}.head h2{margin:0;color:var(--plum);font-size:17px}.ingest{display:grid;grid-template-columns:1fr auto;gap:10px;padding:17px}.ingest input,.controls input,.controls select{height:44px;border:1px solid #d9c9dc;border-radius:10px;padding:0 12px;background:#fff}.helper,.muted{color:var(--muted);font-size:11px}.helper{padding:0 17px 15px}.error{padding:0 17px 15px;color:var(--red);font-size:12px}button{border:0;border-radius:10px;padding:11px 15px;font-weight:800;background:var(--plum);color:#fff;cursor:pointer}.download{background:var(--grad)}.hidden{display:none!important}.loading{display:flex;gap:12px;align-items:center;padding:18px}.spin{width:25px;height:25px;border:3px solid #eadfea;border-top-color:var(--purple);border-radius:50%;animation:s .8s linear infinite}@keyframes s{to{transform:rotate(360deg)}}.cards{display:grid;grid-template-columns:repeat(6,minmax(140px,1fr));gap:12px;margin:18px 0}.card{background:#fff;border:1px solid var(--line);border-radius:15px;padding:14px 15px;box-shadow:0 6px 20px rgba(53,18,63,.05)}.label{font-size:10px;font-weight:850;text-transform:uppercase;color:var(--muted)}.value{font-size:27px;font-weight:880;color:var(--plum);margin-top:5px}.alert .value{color:var(--red)}.grid{display:grid;grid-template-columns:1fr 1fr;gap:16px}.body{padding:15px 17px}.bar{display:grid;grid-template-columns:minmax(170px,1.5fr) 2fr 54px;gap:10px;align-items:center;margin:10px 0;font-size:11px}.track{height:9px;border-radius:999px;background:#f1e7f3;overflow:hidden}.fill{height:100%;background:var(--grad);border-radius:999px}.controls{display:grid;grid-template-columns:2fr 1fr 1fr auto;gap:10px;align-items:center;margin:16px 0}.table{overflow:auto;max-height:68vh;border:1px solid var(--line);border-radius:15px;background:#fff}table{border-collapse:separate;border-spacing:0;min-width:2800px;width:100%;font-size:11px}th{position:sticky;top:0;background:var(--plum);color:#fff;padding:10px;text-align:left}td{padding:9px;vertical-align:top;border-bottom:1px solid #eee4ef;border-right:1px solid #f5eff6;max-width:360px}tr.review td{background:#fff6f4}.badge{display:inline-flex;border-radius:999px;padding:4px 8px;font-size:10px;font-weight:850}.yes{color:var(--red);background:#fee4e2}.no{color:var(--green);background:#dcfce7}.reviewb{color:#fff;background:var(--red)}.pill{display:inline-block;border-radius:8px;padding:5px 7px;background:#f2e6f8;color:var(--plum);font-weight:800}.high{color:var(--green);font-weight:850}.mid{color:var(--amber);font-weight:850}.low{color:var(--red);font-weight:850}footer{text-align:center;margin-top:18px;color:var(--muted);font-size:11px}@media(max-width:1000px){.cards{grid-template-columns:repeat(3,1fr)}.grid{grid-template-columns:1fr}.controls{grid-template-columns:1fr 1fr}}@media(max-width:650px){header,main{padding-left:14px;padding-right:14px}.ingest{grid-template-columns:1fr}.cards{grid-template-columns:1fr 1fr}.controls{grid-template-columns:1fr}}
</style></head><body><div class="rule"></div><header><div class="top"><div class="brand">Grit <span>Financial</span></div><div class="status">Compliance Analytics • Human-governed preliminary triage</div></div><div class="hero"><div class="eyebrow">Consumer Complaint Review</div><h1>Complaint Classification & Regulatory Risk Analytics</h1><p>Analyze a Google Sheet of customer complaints, identify potential UDAAP / Reg E / Reg DD / Reg P / Reg GG indicators, route low-confidence cases to human review, and export an audit-ready Excel workbook.</p></div></header><main>
<div class="notice"><strong>Important:</strong> Results are preliminary compliance triage, not legal determinations. Confidence reflects support for a classification from the available record—not the probability that a violation occurred. Final determinations remain subject to qualified human review.</div>
<section class="panel"><div class="head"><h2>Run a classification</h2><span class="muted">80% review threshold</span></div><div class="ingest"><input id="url" placeholder="https://docs.google.com/spreadsheets/d/.../edit#gid=0"><button id="go">Analyze complaints</button></div><div class="helper">Required: complaint/reason/description column. Recommended: ticket ID, date, customer, product/program, resolution.</div><div id="err" class="error"></div></section>
<section id="loading" class="panel loading hidden"><div class="spin"></div><div><strong>Analyzing complaints…</strong><div class="muted">Reading the sheet, applying regulatory tests, and building the review queue.</div></div></section>
<section id="dash" class="hidden"><div class="head" style="padding:0;border:0;background:transparent;margin-top:24px"><div><div class="eyebrow" style="color:var(--purple)">Current run</div><h2>Compliance analytics dashboard</h2></div><button id="download" class="download">Download .xlsx</button></div><div id="cards" class="cards"></div><div class="grid"><section class="panel"><div class="head"><h2>Classification mix</h2></div><div id="classmix" class="body"></div></section><section class="panel"><div class="head"><h2>Top root causes</h2></div><div id="rootmix" class="body"></div></section></div><div class="controls"><input id="search" placeholder="Search ticket, customer, complaint, rationale…"><select id="cf"><option value="">All classifications</option></select><select id="rf"><option value="">All review statuses</option><option value="Yes">Human review required</option><option value="No">No review required</option></select><span id="count" class="muted"></span></div><div class="table"><table><thead><tr><th>Case ID</th><th>Customer</th><th>Date</th><th>Program / Product</th><th>Complaint</th><th>Issue category</th><th>UDAAP</th><th>Reg E</th><th>Reg DD</th><th>Reg P</th><th>Reg GG</th><th>Classification</th><th>Confidence</th><th>Human review</th><th>Priority</th><th>Rationale</th><th>Missing facts</th><th>Root cause</th></tr></thead><tbody id="tbody"></tbody></table></div></section>
<footer>Grit Financial • Compliance Analytics • Preliminary classification for human review</footer></main><script>
let run=null,rows=[];const $=x=>document.getElementById(x),esc=s=>String(s??'').replace(/[&<>"']/g,m=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[m]));const badge=v=>`<span class="badge ${v==='Yes'?'yes':'no'}">${v}</span>`;function counts(k){let o={};rows.forEach(r=>o[r[k]||'Unknown']=(o[r[k]||'Unknown']||0)+1);return Object.entries(o).sort((a,b)=>b[1]-a[1])}function bars(id,data){let m=data.length?data[0][1]:1;$(id).innerHTML=data.slice(0,8).map(([k,v])=>`<div class="bar"><div>${esc(k)}</div><div class="track"><div class="fill" style="width:${100*v/m}%"></div></div><strong>${v}</strong></div>`).join('')}function cards(s){$('cards').innerHTML=[['Cases',s.total_cases,'Rows analyzed'],['Needs human review',s.needs_human_review,'Confidence below 80%',1],['Average confidence',s.average_confidence+'%','Classification support'],['UDAAP flags',s.udaap_flags,'Potential indicator'],['Reg E flags',s.reg_e_flags,'Potential indicator'],['Unknown',s.unknown,'Below 50% confidence',s.unknown>0]].map(x=>`<div class="card ${x[3]?'alert':''}"><div class="label">${x[0]}</div><div class="value">${x[1]}</div><div class="muted">${x[2]}</div></div>`).join('')}function filter(){let q=$('search').value.toLowerCase(),c=$('cf').value,r=$('rf').value;return rows.filter(x=>(!q||Object.values(x).join(' ').toLowerCase().includes(q))&&(!c||x.primary_classification===c)&&(!r||x.needs_human_review===r))}function table(){let f=filter();$('count').textContent=`${f.length} of ${rows.length} cases`;$('tbody').innerHTML=f.map(r=>`<tr class="${r.needs_human_review==='Yes'?'review':''}"><td><strong>${esc(r.case_id)}</strong></td><td>${esc(r.customer)}</td><td>${esc(r.date)}</td><td>${esc(r.program_product)}</td><td>${esc(r.complaint)}</td><td>${esc(r.issue_category)}</td><td>${badge(r.udaap)}</td><td>${badge(r.reg_e)}</td><td>${badge(r.reg_dd)}</td><td>${badge(r.reg_p)}</td><td>${badge(r.reg_gg)}</td><td><span class="pill">${esc(r.primary_classification)}</span></td><td><span class="${r.confidence>=80?'high':r.confidence>=50?'mid':'low'}">${Number(r.confidence).toFixed(1)}%</span></td><td>${r.needs_human_review==='Yes'?'<span class="badge reviewb">Review</span>':'<span class="badge no">Clear</span>'}</td><td>${esc(r.priority)}</td><td>${esc(r.regulatory_rationale)}</td><td>${esc(r.missing_facts)}</td><td>${esc(r.root_cause_category)}</td></tr>`).join('')}async function go(){$('err').textContent='';$('loading').classList.remove('hidden');$('dash').classList.add('hidden');$('go').disabled=true;try{let res=await fetch('/api/analyze',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({sheet_url:$('url').value.trim()})});let d=await res.json();if(!res.ok)throw new Error(d.detail||'Analysis failed');run=d.run_id;rows=d.results;cards(d.summary);bars('classmix',counts('primary_classification'));bars('rootmix',counts('root_cause_category'));$('cf').innerHTML='<option value="">All classifications</option>'+[...new Set(rows.map(x=>x.primary_classification))].sort().map(v=>`<option>${esc(v)}</option>`).join('');table();$('dash').classList.remove('hidden')}catch(e){$('err').textContent=e.message}finally{$('loading').classList.add('hidden');$('go').disabled=false}}$('go').onclick=go;$('download').onclick=()=>{if(run)location=`/api/runs/${run}/download.xlsx`};['search','cf','rf'].forEach(id=>$(id).addEventListener(id==='search'?'input':'change',table));
</script></body></html>'''
